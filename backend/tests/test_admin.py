"""Admin surface: auth, import from raw notes, the publish gate, and export."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.models import AllergenStatus, ProductStatus
from tests.conftest import ADMIN_HEADERS, make_product


class TestAuth:
    def test_admin_endpoints_reject_a_missing_key(self, client):
        assert client.get("/api/admin/readiness").status_code == 401

    def test_admin_endpoints_reject_a_wrong_key(self, client):
        response = client.get("/api/admin/readiness", headers={"X-Admin-Key": "guess"})
        assert response.status_code == 401

    def test_a_valid_key_is_accepted(self, client):
        assert client.get("/api/admin/readiness", headers=ADMIN_HEADERS).status_code == 200


class TestImport:
    def test_preview_returns_records_without_writing(self, client, raw_products):
        body = client.post(
            "/api/admin/import/preview",
            headers=ADMIN_HEADERS,
            json={"raw_text": raw_products},
        ).json()
        assert len(body["records"]) == 20
        assert body["categories"] == ["Hummus", "Pesto", "Olives", "Sweet Pepper Drops"]
        assert body["can_commit"] is True
        assert client.get("/api/products", params={"include_drafts": True}).json()["total"] == 0

    def test_preview_reports_ambiguities_with_line_numbers(self, client):
        body = client.post(
            "/api/admin/import/preview",
            headers=ADMIN_HEADERS,
            json={"raw_text": "- Orphan Olives\n"},
        ).json()
        assert body["can_commit"] is False
        assert body["ambiguities"][0]["line"] == 1
        assert body["ambiguities"][0]["kind"] == "orphan_product"

    def test_commit_creates_draft_products(self, client, raw_products):
        body = client.post(
            "/api/admin/import/commit",
            headers=ADMIN_HEADERS,
            json={"raw_text": raw_products},
        ).json()
        assert len(body["created_products"]) == 20
        assert len(body["created_categories"]) == 4

        listed = client.get(
            "/api/products", params={"include_drafts": True, "limit": 50}
        ).json()
        assert listed["total"] == 20
        assert all(item["status"] == "draft" for item in listed["items"])

    def test_commit_invents_no_prices_or_allergens(self, client, raw_products):
        """Imported products carry names only. Everything else is the client's to supply."""
        client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        )
        detail = client.get(
            "/api/products/hummus-traditional-hummus", params={"include_drafts": True}
        ).json()
        assert detail["variants"] == []
        assert detail["allergens"] == []
        assert detail["allergen_status"] == "unconfirmed"
        assert detail["description"] is None

    def test_commit_is_blocked_by_ambiguity(self, client):
        response = client.post(
            "/api/admin/import/commit",
            headers=ADMIN_HEADERS,
            json={"raw_text": "Olives\n- House Mix\n    - Extra Garlic\n"},
        )
        assert response.status_code == 409
        assert response.json()["detail"]["ambiguities"][0]["kind"] == "nested_bullet"
        assert client.get("/api/products", params={"include_drafts": True}).json()["total"] == 0

    def test_ambiguity_can_be_overridden_explicitly(self, client):
        response = client.post(
            "/api/admin/import/commit",
            headers=ADMIN_HEADERS,
            json={
                "raw_text": "Olives\n- House Mix\n    - Extra Garlic\n",
                "allow_ambiguous": True,
            },
        )
        assert response.status_code == 200
        assert response.json()["created_products"] == ["House Mix"]

    def test_reimporting_skips_products_that_already_exist(self, client, raw_products):
        client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        )
        second = client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        ).json()
        assert second["created_products"] == []
        assert len(second["skipped_existing"]) == 20

    def test_empty_input_is_refused(self, client):
        response = client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": "   \n  "}
        )
        assert response.status_code == 400


class TestExport:
    def test_export_is_a_real_xlsx_matching_the_catalogue(self, client, raw_products):
        client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        )
        response = client.get("/api/admin/export.xlsx", headers=ADMIN_HEADERS)
        assert response.status_code == 200
        assert response.content[:2] == b"PK"  # a zip container, not a renamed CSV

        sheet = load_workbook(io.BytesIO(response.content))["Products"]
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0] == ("Category", "Product")
        assert len(rows) == 21
        assert rows[1] == ("Hummus", "Traditional Hummus")
        assert ("Olives", "Garlicy Green Olives") in rows
        assert ("Sweet Pepper Drops", "Red") in rows

    def test_export_of_an_empty_catalogue_is_a_404_not_an_empty_file(self, client):
        assert client.get("/api/admin/export.xlsx", headers=ADMIN_HEADERS).status_code == 404


class TestPublishing:
    def test_a_draft_with_no_allergen_data_cannot_be_published(self, client, db_session):
        make_product(
            db_session,
            status=ProductStatus.draft,
            allergen_status=AllergenStatus.unconfirmed,
        )
        response = client.post(
            "/api/admin/products/hummus-traditional-hummus/publish", headers=ADMIN_HEADERS
        )
        assert response.status_code == 409
        blockers = response.json()["detail"]["blockers"]
        assert any("Allergen data is not confirmed" in blocker for blocker in blockers)

    def test_a_ready_product_publishes(self, client, db_session):
        make_product(db_session, status=ProductStatus.draft)
        response = client.post(
            "/api/admin/products/hummus-traditional-hummus/publish", headers=ADMIN_HEADERS
        )
        assert response.status_code == 200
        assert response.json()["status"] == "active"
        assert client.get("/api/products").json()["total"] == 1

    def test_unpublishing_removes_it_from_the_storefront(self, client, db_session):
        make_product(db_session)
        client.post(
            "/api/admin/products/hummus-traditional-hummus/unpublish", headers=ADMIN_HEADERS
        )
        assert client.get("/api/products").json()["total"] == 0

    def test_readiness_lists_what_each_product_still_needs(self, client, raw_products):
        client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        )
        body = client.get("/api/admin/readiness", headers=ADMIN_HEADERS).json()
        assert body["summary"]["total_products"] == 20
        assert body["summary"]["blocked"] == 20
        assert body["summary"]["launch_ready"] is False
        first = body["products"][0]
        assert any("Allergen data is not confirmed" in blocker for blocker in first["blockers"])
        assert any("No sellable variant" in blocker for blocker in first["blockers"])


class TestProductEditing:
    def test_variants_and_allergens_can_be_set(self, client, raw_products):
        client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        )
        response = client.patch(
            "/api/admin/products/hummus-traditional-hummus",
            headers=ADMIN_HEADERS,
            json={
                "temperature": "chilled",
                "is_vegan": True,
                "is_vegetarian": True,
                "allergen_status": "confirmed",
                "allergens": [{"allergen": "sesame", "presence": "contains"}],
                "variants": [
                    {"label": "200g", "price_pence": 450, "weight_grams": 200, "stock": 24}
                ],
            },
        )
        assert response.status_code == 200
        body = response.json()
        assert body["variants"][0]["sku"] == "HUM-TRADITIO-200G"
        assert body["allergens"][0]["allergen"] == "sesame"

        published = client.post(
            "/api/admin/products/hummus-traditional-hummus/publish", headers=ADMIN_HEADERS
        )
        assert published.status_code == 200

    def test_an_unregulated_allergen_name_is_rejected(self, client, db_session):
        make_product(db_session)
        response = client.patch(
            "/api/admin/products/hummus-traditional-hummus",
            headers=ADMIN_HEADERS,
            json={"allergens": [{"allergen": "chickpeas", "presence": "contains"}]},
        )
        assert response.status_code == 400
        assert "Not a regulated allergen" in response.json()["detail"]

    def test_the_fourteen_regulated_allergens_are_listed(self, client):
        body = client.get("/api/admin/allergens", headers=ADMIN_HEADERS).json()
        assert len(body) == 14
        assert "milk" in body and "sulphur_dioxide" in body

    def test_editing_never_renames_a_product(self, client, raw_products):
        """The name comes from the source notes and is not editable through this API."""
        client.post(
            "/api/admin/import/commit", headers=ADMIN_HEADERS, json={"raw_text": raw_products}
        )
        client.patch(
            "/api/admin/products/olives-garlicy-green-olives",
            headers=ADMIN_HEADERS,
            json={"description": "House favourite"},
        )
        body = client.get(
            "/api/products/olives-garlicy-green-olives", params={"include_drafts": True}
        ).json()
        assert body["name"] == "Garlicy Green Olives"

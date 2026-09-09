"""Export integrity. A silent off-by-one in an export is worse than a loud failure."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from app.services.xlsx import ExportError, build_workbook_bytes, verify_workbook_bytes

ROWS = [("Hummus", "Traditional Hummus"), ("Olives", "Garlicy Green Olives")]


def test_workbook_has_the_expected_shape():
    sheet = load_workbook(io.BytesIO(build_workbook_bytes(ROWS)))["Products"]
    assert list(sheet.iter_rows(values_only=True)) == [("Category", "Product"), *ROWS]
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold is True


def test_it_is_a_real_xlsx_not_a_renamed_csv():
    assert build_workbook_bytes(ROWS)[:2] == b"PK"


def test_names_are_written_verbatim():
    sheet = load_workbook(io.BytesIO(build_workbook_bytes(ROWS)))["Products"]
    assert sheet["B3"].value == "Garlicy Green Olives"


def test_an_empty_export_fails_loudly():
    with pytest.raises(ExportError, match="No rows to export"):
        build_workbook_bytes([])


def test_verification_catches_a_row_count_mismatch():
    data = build_workbook_bytes(ROWS)
    with pytest.raises(ExportError, match="Row count mismatch"):
        verify_workbook_bytes(data, ROWS[:1])


def test_verification_catches_a_changed_cell():
    data = build_workbook_bytes(ROWS)
    with pytest.raises(ExportError, match="mismatch"):
        verify_workbook_bytes(data, [ROWS[0], ("Olives", "Garlicky Green Olives")])


def test_verification_catches_a_missing_sheet():
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(ExportError, match="missing the 'Products' sheet"):
        verify_workbook_bytes(buffer.getvalue(), ROWS)

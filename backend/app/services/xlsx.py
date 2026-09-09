"""Export `Category | Product` rows to a genuine .xlsx workbook.

Written with openpyxl — never a CSV renamed. The workbook is read back and
checked cell by cell before it is handed over: a silent off-by-one in an export
is worse than a loud failure.
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADER = ("Category", "Product")
SHEET_NAME = "Products"


class ExportError(RuntimeError):
    pass


def build_workbook_bytes(rows: list[tuple[str, str]]) -> bytes:
    if not rows:
        raise ExportError("No rows to export; nothing was written.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(list(HEADER))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for category, product in rows:
        sheet.append([category, product])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 42

    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    verify_workbook_bytes(data, rows)
    return data


def verify_workbook_bytes(data: bytes, rows: list[tuple[str, str]]) -> None:
    """Re-open the workbook and assert it matches the rows that were sent."""
    workbook = load_workbook(io.BytesIO(data))
    if SHEET_NAME not in workbook.sheetnames:
        raise ExportError(f"Workbook is missing the {SHEET_NAME!r} sheet.")
    sheet = workbook[SHEET_NAME]
    written = list(sheet.iter_rows(values_only=True))

    if written[0] != HEADER:
        raise ExportError(f"Bad header row: {written[0]!r}")
    if len(written) - 1 != len(rows):
        raise ExportError(
            f"Row count mismatch: {len(written) - 1} in file vs {len(rows)} sent."
        )
    for index, (actual, expected) in enumerate(zip(written[1:], rows), start=2):
        if actual != expected:
            raise ExportError(f"Row {index} mismatch: {actual!r} != {expected!r}")
